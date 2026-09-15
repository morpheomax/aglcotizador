from datetime import date
from io import BytesIO

import openpyxl
import pandas as pd

from src.exporters.excel_exporter import export_quote_to_excel
from src.models.project import ProjectInfo


def test_export_quote_to_excel_has_required_sheets() -> None:
    project = ProjectInfo("Cliente", "Proyecto", "Chile", "Santiago", 0, "USD", None, date(2026, 7, 6))
    rooms_df = pd.DataFrame([{"name": "Sala 1", "agent_type": "FM-200"}])
    results_df = pd.DataFrame([{"room_name": "Sala 1", "agent_required_rounded_kg": 65.0}])
    cylinders_df = pd.DataFrame([{"room_name": "Sala 1", "cylinder_code": "FM200-106L"}])
    bom_detail_df = pd.DataFrame([{"code": "FM200-AGENT", "product": "Agente", "quantity": 65.0}])
    bom_consolidated_df = pd.DataFrame([{"code": "FM200-AGENT", "product": "Agente", "quantity": 65.0}])

    excel_bytes = export_quote_to_excel(
        project,
        rooms_df,
        results_df,
        cylinders_df,
        bom_detail_df,
        bom_consolidated_df,
        warnings=["Advertencia de prueba"],
    )
    workbook = openpyxl.load_workbook(BytesIO(excel_bytes), read_only=True)

    assert set(workbook.sheetnames) == {
        "Resumen",
        "Salas",
        "Resultados_Tecnicos",
        "Cilindros",
        "BOM_Detalle",
        "BOM_Consolidado",
        "Advertencias",
    }
    assert workbook["Advertencias"].max_row == 3
