"""Excel export utilities."""

from __future__ import annotations

from dataclasses import asdict, is_dataclass
from io import BytesIO
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


COMPLIANCE_NOTE = (
    "Cálculo preliminar para cotización. Debe ser validado mediante cálculo "
    "hidráulico/software certificado y revisión técnica antes de emisión final "
    "para instalación."
)


def export_quote_to_excel(
    project_info: Any,
    rooms_df: pd.DataFrame,
    results_df: pd.DataFrame,
    cylinders_df: pd.DataFrame,
    bom_detail_df: pd.DataFrame,
    bom_consolidated_df: pd.DataFrame,
    warnings: list[str] | None = None,
) -> bytes:
    """Build an in-memory Excel quote workbook."""
    output = BytesIO()
    sheets = {
        "Resumen": _project_to_dataframe(project_info),
        "Salas": rooms_df,
        "Resultados_Tecnicos": results_df,
        "Cilindros": cylinders_df,
        "BOM_Detalle": bom_detail_df,
        "BOM_Consolidado": bom_consolidated_df,
        "Advertencias": _warnings_to_dataframe(warnings or []),
    }

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, df in sheets.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)
        _format_workbook(writer, sheets)

    return output.getvalue()


def _project_to_dataframe(project_info: Any) -> pd.DataFrame:
    if project_info is None:
        return pd.DataFrame()
    if is_dataclass(project_info):
        data = asdict(project_info)
    elif isinstance(project_info, dict):
        data = project_info
    else:
        data = dict(project_info)
    return pd.DataFrame([data])


def _warnings_to_dataframe(warnings: list[str]) -> pd.DataFrame:
    rows = [{"tipo": "Compliance", "mensaje": COMPLIANCE_NOTE}]
    rows.extend({"tipo": "Advertencia", "mensaje": message} for message in warnings)
    return pd.DataFrame(rows)


def _format_workbook(writer: pd.ExcelWriter, sheets: dict[str, pd.DataFrame]) -> None:
    workbook = writer.book
    header_fill = PatternFill(fill_type="solid", fgColor="1E3A5F")
    header_font = Font(color="FFFFFF", bold=True)
    body_alignment = Alignment(wrap_text=True, vertical="top")

    for sheet_name, df in sheets.items():
        worksheet = workbook[sheet_name]
        worksheet.freeze_panes = "A2"
        for column_index, column_name in enumerate(df.columns, start=1):
            cell = worksheet.cell(row=1, column=column_index)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = body_alignment
            width = min(max(len(str(column_name)) + 4, 14), 36)
            worksheet.column_dimensions[cell.column_letter].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = body_alignment
        if len(df.columns) > 0:
            worksheet.auto_filter.ref = worksheet.dimensions
